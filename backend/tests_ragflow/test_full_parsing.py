"""Generic CPU-format and pinned-symbol parity tests; no quality-set fixtures."""
import ast
from io import BytesIO
from pathlib import Path
import zipfile

import pytest
from ragflow_derived.contracts import EngineError
from ragflow_derived.parsing import parse
from ragflow_derived.upstream import merge, naive_markdown
from .fixtures import TokenizerDouble


def pinned_symbols(relative, names, namespace):
    source = Path(__file__).resolve().parents[2] / '.codex_ragflow_upstream' / relative
    tree = ast.parse(source.read_text(encoding='utf-8'))
    selected = [node for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef))
                and node.name in names]
    assert {n.name for n in selected} == set(names)
    exec(compile(ast.Module(body=selected, type_ignores=[]), relative, 'exec'), namespace)
    return namespace


@pytest.fixture(autouse=True)
def deterministic_tokens(monkeypatch):
    monkeypatch.setattr(merge, 'num_tokens_from_string', lambda text: len(text.split()))
    monkeypatch.setattr(naive_markdown, 'num_tokens_from_string', lambda text: len(text.split()))


@pytest.mark.parametrize('delimiter', ['\n', '', '`END`'])
@pytest.mark.parametrize('budget', [1, 8, 512])
def test_naive_merge_actual_pinned_symbols(delimiter, budget):
    names = ['_compute_overlap_prefix', '_merge_paragraph_groups', '_reconstruct_text_chunk',
             '_apply_overlap_unconditional', 'naive_merge']
    original = pinned_symbols('rag/nlp/__init__.py', names, dict(vars(merge)))
    sections = [('First line\nSecond line END next fragment.', '@@1\t2##'), ('Final text.', '')]
    assert merge.naive_merge(sections, budget, delimiter) == original['naive_merge'](sections, budget, delimiter)


@pytest.mark.parametrize('context_size', [0, 5, 20])
def test_docx_merger_preserves_pinned_behavior_including_table_placement(context_size):
    from copy import deepcopy
    names = ['_build_cks', '_add_context', '_merge_cks', 'naive_merge_docx']
    original = pinned_symbols('rag/nlp/__init__.py', names, dict(vars(merge)))
    rows = [('Before table', None, ''), ('', None, '<table><tr><td>2 kg</td></tr></table>'),
            ('After table', None, '')]
    assert merge.naive_merge_docx(deepcopy(rows), 8, table_context_size=context_size) == \
        original['naive_merge_docx'](deepcopy(rows), 8, table_context_size=context_size)


def test_docx_dispatch_uses_real_parser_and_merge_with_hierarchy():
    from docx import Document
    from ragflow_derived.upstream.naive_docx import Docx
    doc = Document()
    doc.add_heading('Laboratory', 1)
    doc.add_paragraph('Before table')
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip([c for r in table.rows for c in r.cells], ['Item', 'Count', 'Glass', '4']):
        cell.text = text
    doc.add_paragraph('After table')
    stream = BytesIO()
    doc.save(stream)
    raw = stream.getvalue()
    expected, _ = merge.naive_merge_docx(Docx()('manual.docx', raw), 512)
    actual = parse(raw, 'docx', tokenizer=TokenizerDouble(), filename='manual.docx')
    assert [p['text'] for p in actual] == [p['text'] for p in expected if p['text'].strip()]
    assert any('Laboratory' in p['text'] and '<table' in p['text'] for p in actual)
    assert all(p['parser_metadata']['parser'].endswith('naive_merge_docx') for p in actual)


@pytest.mark.parametrize('kind,raw', [('json', b'{"zero":0,"active":false,"title":"Notes"}'),
                                    ('jsonl', b'{"row":1}\n{"row":2}\n')])
def test_json_parser_dispatch_exact(kind, raw):
    from ragflow_derived.upstream.json_parser import RAGFlowJsonParser
    actual = parse(raw, kind, tokenizer=TokenizerDouble())
    assert [p['text'] for p in actual] == RAGFlowJsonParser(512)(raw)
    assert all(p['kind'] == 'structured' for p in actual)


@pytest.mark.parametrize('kind', ['xlsx', 'csv'])
def test_spreadsheet_tables_values_and_positions(kind):
    from openpyxl import Workbook
    from ragflow_derived.upstream.excel_parser import RAGFlowExcelParser
    if kind == 'csv':
        raw = b'Unit,Count\nkg,2\nml,7\n'
    else:
        book = Workbook()
        book.active.title = 'Measures'
        for row in [('Unit', 'Count'), ('kg', 2), ('ml', 7)]:
            book.active.append(row)
        stream = BytesIO()
        book.save(stream)
        raw = stream.getvalue()
    expected = RAGFlowExcelParser().html(raw)
    actual = parse(raw, kind, tokenizer=TokenizerDouble())
    assert [p['text'] for p in actual] == [t for t, _ in expected]
    assert [p['parser_metadata']['sheet_position'] for p in actual] == [list(pos) for _, pos in expected]
    assert all(p['kind'] == 'table' and '<table' in p['text'] for p in actual)


def test_ppt_preserves_slide_order():
    from pptx import Presentation
    p = Presentation()
    for name in ['First slide', 'Second slide']:
        slide = p.slides.add_slide(p.slide_layouts[1])
        slide.shapes.title.text = name
    stream = BytesIO()
    p.save(stream)
    actual = parse(stream.getvalue(), 'pptx', tokenizer=TokenizerDouble())
    assert [a['parser_metadata']['page'] for a in actual] == [1, 2]
    assert 'First slide' in actual[0]['text'] and 'Second slide' in actual[1]['text']


def test_markdown_uses_pinned_header_merge_not_custom_atomic_block_grouping():
    raw = '# Instrument\n\nUsage procedure.\n\n| Unit | Count |\n| --- | --- |\n| kg | 2 |\n'
    sections, _, images = naive_markdown.Markdown(4)('manual.md', raw.encode(), separate_tables=False,
                                                    delimiter='\n', return_section_images=True)
    chunks, _ = naive_markdown.merge_markdown_sections(sections, images, 4)
    actual = parse(raw, 'md', tokenizer=TokenizerDouble(), chunk_tokens=4)
    assert [a['text'] for a in actual] == chunks
    assert any('<table>' in a['text'] and 'kg' in a['text'] and '2' in a['text'] for a in actual)
    assert actual[0]['text'].startswith('# Instrument\nUsage procedure.')


def test_markdown_external_references_never_read_network_or_local_files():
    raw = '![private](http://169.254.169.254/latest/meta-data)\n![local](C:/Windows/win.ini)\nText.'
    actual = parse(raw, 'md', tokenizer=TokenizerDouble())
    assert actual
    assert not any('font' in a['text'] for a in actual)


def test_zip_expansion_limit_checked_before_any_parser():
    stream = BytesIO()
    with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('content.xml', 'A' * 100000)
    with pytest.raises(EngineError, match='PARSER_FAILED'):
        parse(stream.getvalue(), 'docx', tokenizer=TokenizerDouble(), max_bytes=2000)


@pytest.mark.parametrize('kind', ['docx', 'xlsx', 'json', 'pptx', 'epub', 'pdf'])
def test_malformed_documents_fail_without_silent_text_fallback(kind):
    with pytest.raises(EngineError, match='PARSER_FAILED'):
        parse(b'not a valid package', kind, tokenizer=TokenizerDouble())
